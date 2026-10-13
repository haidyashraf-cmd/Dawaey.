from __future__ import annotations

import csv
import hashlib
import io
import hmac
import json
import math
import os
import re
import secrets
import sqlite3
import sys
import threading
import time
import urllib.parse
from contextlib import closing
from http.cookies import SimpleCookie
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.error import URLError
from urllib.parse import parse_qs, urlencode, urlsplit
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DATA_PATH = ROOT / "data" / "dawaey-data.json"
DATABASE_PATH = Path(os.environ.get("DAWAEY_DATABASE", ROOT / "data" / "dawaey-users.sqlite3"))
SESSION_COOKIE = "dawaey_session"
SESSION_SECONDS = 60 * 60 * 24 * 14
OAUTH_STATE_COOKIE = "dawaey_google_state"
OAUTH_STATE_SECONDS = 600
PBKDF2_ITERATIONS = 310_000
GEOCODE_LOCK = threading.Lock()
LAST_GEOCODE_REQUEST = 0.0
TELEGRAM_OFFSET = 0
TELEGRAM_THREAD_STARTED = False


def connect_database() -> sqlite3.Connection:
    DATABASE_PATH.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DATABASE_PATH, timeout=10)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def initialize_database() -> None:
    with closing(connect_database()) as connection, connection:
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS accounts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                role TEXT NOT NULL CHECK(role IN ('patient', 'pharmacy')),
                full_name TEXT NOT NULL,
                contact TEXT NOT NULL,
                contact_key TEXT NOT NULL UNIQUE,
                password_salt TEXT NOT NULL,
                password_hash TEXT NOT NULL,
                created_at INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS patient_profiles (
                account_id INTEGER PRIMARY KEY REFERENCES accounts(id) ON DELETE CASCADE,
                governorate TEXT NOT NULL,
                district TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS pharmacy_applications (
                account_id INTEGER PRIMARY KEY REFERENCES accounts(id) ON DELETE CASCADE,
                pharmacy_name TEXT NOT NULL,
                pharmacist_name TEXT NOT NULL,
                license_number TEXT NOT NULL UNIQUE,
                address TEXT NOT NULL,
                district TEXT NOT NULL,
                opening_hours TEXT NOT NULL,
                whatsapp TEXT NOT NULL,
                latitude REAL,
                longitude REAL,
                location_url TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL DEFAULT 'pending' CHECK(status IN ('pending', 'approved', 'rejected'))
            );
            CREATE TABLE IF NOT EXISTS sessions (
                token_hash TEXT PRIMARY KEY,
                account_id INTEGER NOT NULL REFERENCES accounts(id) ON DELETE CASCADE,
                expires_at INTEGER NOT NULL
            );
            CREATE INDEX IF NOT EXISTS sessions_expiry ON sessions(expires_at);
            CREATE TABLE IF NOT EXISTS donation_requests (
                id TEXT PRIMARY KEY,
                account_id INTEGER REFERENCES accounts(id) ON DELETE SET NULL,
                medicine TEXT NOT NULL,
                area TEXT NOT NULL,
                quantity TEXT NOT NULL,
                latitude REAL,
                longitude REAL,
                status TEXT NOT NULL DEFAULT 'pending' CHECK(status IN ('pending', 'accepted', 'completed', 'rejected')),
                created_at INTEGER NOT NULL,
                updated_at INTEGER NOT NULL
            );
            CREATE INDEX IF NOT EXISTS donation_status ON donation_requests(status);
            CREATE INDEX IF NOT EXISTS donation_created ON donation_requests(created_at DESC);
            CREATE TABLE IF NOT EXISTS patient_saved_medicines (
                account_id INTEGER NOT NULL REFERENCES accounts(id) ON DELETE CASCADE,
                medicine_key TEXT NOT NULL,
                created_at INTEGER NOT NULL,
                PRIMARY KEY (account_id, medicine_key)
            );
            CREATE TABLE IF NOT EXISTS notifications (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                account_id INTEGER NOT NULL REFERENCES accounts(id) ON DELETE CASCADE,
                kind TEXT NOT NULL,
                title TEXT NOT NULL,
                message TEXT NOT NULL,
                related_id TEXT,
                is_read INTEGER NOT NULL DEFAULT 0,
                created_at INTEGER NOT NULL
            );
            CREATE INDEX IF NOT EXISTS notifications_account ON notifications(account_id, is_read, created_at DESC);
            CREATE TABLE IF NOT EXISTS pharmacy_data_files (
                account_id INTEGER PRIMARY KEY REFERENCES accounts(id) ON DELETE CASCADE,
                filename TEXT NOT NULL,
                inventory_json TEXT NOT NULL DEFAULT '[]',
                supply_json TEXT NOT NULL DEFAULT '[]',
                pharmacies_json TEXT NOT NULL DEFAULT '[]',
                uploaded_at INTEGER NOT NULL
            );
            """
        )
        columns = {row["name"] for row in connection.execute("PRAGMA table_info(donation_requests)").fetchall()}
        if "national_id" not in columns:
            connection.execute("ALTER TABLE donation_requests ADD COLUMN national_id TEXT NOT NULL DEFAULT ''")
        if "unit" not in columns:
            connection.execute("ALTER TABLE donation_requests ADD COLUMN unit TEXT NOT NULL DEFAULT 'غير محددة'")
        donation_columns = {row["name"] for row in connection.execute("PRAGMA table_info(donation_requests)").fetchall()}
        if "latitude" not in donation_columns:
            connection.execute("ALTER TABLE donation_requests ADD COLUMN latitude REAL")
        if "longitude" not in donation_columns:
            connection.execute("ALTER TABLE donation_requests ADD COLUMN longitude REAL")
        pharmacy_columns = {row["name"] for row in connection.execute("PRAGMA table_info(pharmacy_applications)").fetchall()}
        if "latitude" not in pharmacy_columns:
            connection.execute("ALTER TABLE pharmacy_applications ADD COLUMN latitude REAL")
        if "longitude" not in pharmacy_columns:
            connection.execute("ALTER TABLE pharmacy_applications ADD COLUMN longitude REAL")
        if "location_url" not in pharmacy_columns:
            connection.execute("ALTER TABLE pharmacy_applications ADD COLUMN location_url TEXT NOT NULL DEFAULT ''")


def normalize_contact(value: object) -> tuple[str, str]:
    contact = str(value or "").strip()
    if "@" in contact:
        normalized = contact.casefold()
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", normalized):
            raise ValueError("اكتب بريد إلكتروني صحيح أو رقم موبايل.")
        return contact, normalized

    digits = re.sub(r"\D", "", contact)
    if not 8 <= len(digits) <= 15:
        raise ValueError("اكتب رقم موبايل صحيح.")
    return contact, digits


def password_digest(password: str, salt: bytes) -> bytes:
    return hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, PBKDF2_ITERATIONS)


def external_json_request(url: str, payload: dict, headers: dict | None = None) -> dict:
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    request = Request(url, data=body, headers={"Content-Type": "application/json", **(headers or {})}, method="POST")
    with urlopen(request, timeout=12) as response:
        return json.loads(response.read().decode("utf-8"))


def update_pharmacy_application(account_id: int, status: str) -> bool:
    if status not in {"approved", "rejected"}:
        return False
    with closing(connect_database()) as connection, connection:
        cursor = connection.execute("UPDATE pharmacy_applications SET status = ? WHERE account_id = ?", (status, account_id))
        return cursor.rowcount == 1


def send_telegram_notification(account_id: int, details: dict) -> None:
    token = os.environ.get("TELEGRAM_BOT_TOKEN", "").strip()
    chat_id = os.environ.get("TELEGRAM_CHAT_ID", "").strip()
    missing = [name for name, value in (("TELEGRAM_BOT_TOKEN", token), ("TELEGRAM_CHAT_ID", chat_id)) if not value]
    if missing:
        raise RuntimeError(f"Telegram configuration is missing: {', '.join(missing)}")
    location = f"\nرابط الموقع: {details['locationUrl']}" if details.get("locationUrl") else ""
    text = ("طلب تسجيل صيدلية جديد\n\n"
            f"الصيدلية: {details['pharmacyName']}\n"
            f"المسؤول: {details['pharmacistName']}\n"
            f"الهاتف: {details['contact']}\n"
            f"المنطقة: {details['district']}\n"
            f"الترخيص: {details['licenseNumber']}\n"
            f"العنوان: {details['address']}" + location)
    payload = {"chat_id": chat_id, "text": text, "reply_markup": {"inline_keyboard": [[
        {"text": "قبول", "callback_data": f"pharmacy:approve:{account_id}"},
        {"text": "رفض", "callback_data": f"pharmacy:reject:{account_id}"},
    ]]}}
    external_json_request(f"https://api.telegram.org/bot{token}/sendMessage", payload)


def send_whatsapp_notification(account_id: int, details: dict) -> None:
    token = os.environ.get("WHATSAPP_ACCESS_TOKEN", "").strip()
    phone_id = os.environ.get("WHATSAPP_PHONE_NUMBER_ID", "").strip()
    recipient = os.environ.get("WHATSAPP_RECIPIENT_PHONE", "").strip()
    if not token or not phone_id or not recipient:
        return
    recipient = re.sub(r"\D", "", recipient)
    if recipient.startswith("0"):
        recipient = "20" + recipient[1:]
    payload = {"messaging_product": "whatsapp", "to": recipient, "type": "interactive", "interactive": {
        "type": "button", "body": {"text": f"طلب تسجيل صيدلية جديد: {details['pharmacyName']}\nالمسؤول: {details['pharmacistName']}\nالهاتف: {details['contact']}\nالمنطقة: {details['district']}\nرابط الموقع: {details.get('locationUrl') or 'غير متاح'}"},
        "action": {"buttons": [
            {"type": "reply", "reply": {"id": f"pharmacy:approve:{account_id}", "title": "قبول"}},
            {"type": "reply", "reply": {"id": f"pharmacy:reject:{account_id}", "title": "رفض"}},
        ]},
    }}
    external_json_request(f"https://graph.facebook.com/v23.0/{phone_id}/messages", payload, {"Authorization": f"Bearer {token}"})


def notify_pharmacy_application(account_id: int, details: dict) -> bool:
    telegram_sent = False
    for sender in (send_telegram_notification, send_whatsapp_notification):
        try:
            sender(account_id, details)
            if sender is send_telegram_notification:
                telegram_sent = True
        except Exception as error:
            print(f"Notification failed ({sender.__name__}): {error}", file=sys.stderr)
    return telegram_sent


def handle_approval_callback(callback_data: str) -> str | None:
    match = re.fullmatch(r"pharmacy:(approve|reject):(\d+)", callback_data or "")
    if not match:
        return None
    status = "approved" if match.group(1) == "approve" else "rejected"
    account_id = int(match.group(2))
    return status if update_pharmacy_application(account_id, status) else None


def delete_telegram_webhook(token: str) -> None:
    """Allow long polling to recover if an old Telegram webhook is still set."""
    try:
        result = external_json_request(f"https://api.telegram.org/bot{token}/deleteWebhook", {"drop_pending_updates": False})
        if not result.get("ok"):
            print(f"Telegram deleteWebhook failed: {result}", file=sys.stderr)
    except Exception as error:
        print(f"Telegram deleteWebhook request failed: {error}", file=sys.stderr)


def telegram_poll_loop() -> None:
    global TELEGRAM_OFFSET
    token = os.environ.get("TELEGRAM_BOT_TOKEN", "").strip()
    if not token:
        return
    webhook_reset = False
    while True:
        try:
            query = urllib.parse.urlencode({"timeout": 20, "offset": TELEGRAM_OFFSET})
            with urlopen(f"https://api.telegram.org/bot{token}/getUpdates?{query}", timeout=30) as response:
                result = json.loads(response.read().decode("utf-8"))
            if not result.get("ok", True):
                raise RuntimeError(f"Telegram getUpdates failed: {result}")
            for update in result.get("result", []):
                TELEGRAM_OFFSET = max(TELEGRAM_OFFSET, int(update.get("update_id", 0)) + 1)
                callback = update.get("callback_query") or {}
                data = callback.get("data", "")
                try:
                    status = handle_approval_callback(data)
                except Exception as error:
                    status = None
                    print(f"Telegram approval callback failed for {data!r}: {error}", file=sys.stderr)
                callback_id = callback.get("id")
                if callback_id:
                    callback_text = ("تم اعتماد الطلب" if status == "approved" else "تم رفض الطلب") if status else "الطلب غير موجود أو تم التعامل معه مسبقًا"
                    try:
                        external_json_request(f"https://api.telegram.org/bot{token}/answerCallbackQuery", {"callback_query_id": callback_id, "text": callback_text, "show_alert": not bool(status)})
                    except Exception as error:
                        print(f"Telegram callback acknowledgement failed: {error}", file=sys.stderr)
                if status:
                    message = callback.get("message") or {}
                    chat_id = (message.get("chat") or {}).get("id")
                    message_id = message.get("message_id")
                    if chat_id and message_id:
                        try:
                            external_json_request(f"https://api.telegram.org/bot{token}/editMessageText", {"chat_id": chat_id, "message_id": message_id, "text": f"تم {('قبول' if status == 'approved' else 'رفض')} طلب الصيدلية رقم {data.rsplit(':', 1)[-1]}."})
                        except Exception as error:
                            print(f"Telegram message update failed: {error}", file=sys.stderr)
        except Exception as error:
            print(f"Telegram polling failed: {error}", file=sys.stderr)
            if not webhook_reset and "409" in str(error):
                delete_telegram_webhook(token)
                webhook_reset = True
            time.sleep(5)


def start_telegram_polling() -> None:
    global TELEGRAM_THREAD_STARTED
    if TELEGRAM_THREAD_STARTED or not os.environ.get("TELEGRAM_BOT_TOKEN"):
        return
    TELEGRAM_THREAD_STARTED = True
    threading.Thread(target=telegram_poll_loop, name="telegram-approval-poller", daemon=True).start()


def create_session(connection: sqlite3.Connection, account_id: int) -> str:
    token = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(token.encode("ascii")).hexdigest()
    connection.execute(
        "INSERT INTO sessions(token_hash, account_id, expires_at) VALUES (?, ?, ?)",
        (token_hash, account_id, int(time.time()) + SESSION_SECONDS),
    )
    return token


def public_account(connection: sqlite3.Connection, account: sqlite3.Row) -> dict:
    result = {
        "id": account["id"],
        "role": account["role"],
        "name": account["full_name"],
        "contact": account["contact"],
    }
    if account["role"] == "pharmacy":
        application = connection.execute(
            "SELECT pharmacy_name, status FROM pharmacy_applications WHERE account_id = ?",
            (account["id"],),
        ).fetchone()
        if application:
            result["pharmacyName"] = application["pharmacy_name"]
            result["status"] = application["status"]
    return result


def oauth_state_cookie(state: str) -> str:
    secure = os.environ.get("DAWAEY_SECURE_COOKIE") == "1"
    suffix = "; Secure" if secure else ""
    return f"{OAUTH_STATE_COOKIE}={state}; HttpOnly; SameSite=Lax; Path=/; Max-Age={OAUTH_STATE_SECONDS}{suffix}"


def expired_oauth_state_cookie() -> str:
    secure = os.environ.get("DAWAEY_SECURE_COOKIE") == "1"
    suffix = "; Secure" if secure else ""
    return f"{OAUTH_STATE_COOKIE}=; HttpOnly; SameSite=Lax; Path=/; Max-Age=0{suffix}"


def get_cookie_value(handler, name: str) -> str | None:
    cookie = SimpleCookie()
    try:
        cookie.load(handler.headers.get("Cookie", ""))
    except Exception:
        return None
    morsel = cookie.get(name)
    return morsel.value if morsel else None


def google_redirect_uri(handler) -> str:
    base = os.environ.get("DAWAEY_PUBLIC_URL", "").strip().rstrip("/")
    if not base:
        scheme = handler.headers.get("X-Forwarded-Proto", "https")
        base = f"{scheme}://{handler.headers.get('Host', 'localhost')}"
    return f"{base}/auth/google/callback"


def google_token_exchange(code: str, redirect_uri: str, client_id: str, client_secret: str) -> dict:
    payload = urlencode({"code": code, "client_id": client_id, "client_secret": client_secret, "redirect_uri": redirect_uri, "grant_type": "authorization_code"}).encode("utf-8")
    request = Request("https://oauth2.googleapis.com/token", data=payload, headers={"Content-Type": "application/x-www-form-urlencoded"}, method="POST")
    with urlopen(request, timeout=12) as response:
        return json.loads(response.read().decode("utf-8"))


def google_user_profile(access_token: str) -> dict:
    request = Request("https://openidconnect.googleapis.com/v1/userinfo", headers={"Authorization": f"Bearer {access_token}"})
    with urlopen(request, timeout=12) as response:
        return json.loads(response.read().decode("utf-8"))


def uploaded_json_value(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def clean_uploaded_rows(rows):
    aliases = {
        "اسم الصنف": "اسم الدواء", "الصنف": "اسم الدواء", "الدواء": "اسم الدواء", "اسم المنتج": "اسم الدواء",
        "رقم الصنف": "الكود", "كود الصنف": "الكود", "رقم الصنف/الكود": "الكود",
        "الفئة": "طبيعة الدواء", "التصنيف": "طبيعة الدواء", "نوع الدواء": "طبيعة الدواء",
        "الوحدة": "وحدة القياس", "نوع العبوة": "وحدة القياس", "العبوة": "وحدة القياس",
        "الكمية": "الكمية الحالية", "الرصيد": "الكمية الحالية", "الكمية المتاحة": "الكمية الحالية",
        "حد الطلب": "الحد الأدنى", "الحد الادنى": "الحد الأدنى", "الحد الأدنى للمخزون": "الحد الأدنى",
        "توريد مقترح": "كمية التوريد المقترحة", "الكمية المطلوبة": "كمية التوريد المقترحة",
        "الحالة": "حالة المخزون", "حالة الصنف": "حالة المخزون",
        "اسم الفرع": "اسم الصيدلية", "اسم الصيدليه": "اسم الصيدلية",
        "العنوان": "المنطقة / العنوان المتوقع", "المنطقة": "المنطقة / العنوان المتوقع",
        "الهاتف": "رقم التليفون / الخط الساخن", "رقم الهاتف": "رقم التليفون / الخط الساخن", "التليفون": "رقم التليفون / الخط الساخن",
    }
    clean=[]
    for row in rows[:10000]:
        if not isinstance(row, dict): continue
        normalized={}
        for key,value in row.items():
            key=str(key).strip()
            if not key: continue
            normalized[key]=uploaded_json_value(value)
            canonical=aliases.get(key)
            if canonical and not normalized.get(canonical): normalized[canonical]=uploaded_json_value(value)
        if any(value != "" for value in normalized.values()): clean.append(normalized)
    return clean


def find_header_row(values):
    keywords=("اسم الدواء","اسم الصنف","الصنف","الدواء","الكمية","الرصيد","الكود","اسم الصيدلية","الهاتف","التليفون")
    best_index=0; best_score=-1
    for index,row in enumerate(values[:15]):
        cells=[str(value).strip() for value in row if value is not None and str(value).strip()]
        score=sum(1 for cell in cells for keyword in keywords if keyword in cell)
        if score > best_score or (score == best_score and len(cells)>0 and len(cells)>len(values[best_index])):
            best_index=index; best_score=score
    return best_index


def classify_uploaded_sheet(name, rows):
    name_text = str(name or "").lower()
    headers = " ".join(str(key) for row in rows[:3] for key in row.keys())
    if any(token in name_text or token in headers for token in ("توريد", "كمية التوريد", "الحد الأدنى")):
        return "supply"
    if any(token in name_text or token in headers for token in ("صيدال", "pharmacy", "اسم الصيدلية", "رقم التليفون")):
        return "pharmacies"
    return "inventory"


def parse_uploaded_rows(content, filename):
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as error:
            raise ValueError("تعذر قراءة ملف Excel. تأكد أنه ملف XLSX سليم.") from error
        parsed = {}
        for sheet in workbook.worksheets:
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                continue
            header_index = find_header_row(values)
            headers = [str(value).strip() if value is not None else "" for value in values[header_index]]
            rows = []
            for values_row in values[header_index + 1:]:
                row = {headers[index]: uploaded_json_value(value) for index, value in enumerate(values_row) if index < len(headers) and headers[index]}
                if any(value != "" for value in row.values()):
                    rows.append(row)
            if rows:
                parsed[classify_uploaded_sheet(sheet.title, rows)] = clean_uploaded_rows(rows)
    elif suffix in {".csv", ".txt"}:
        text = None
        for encoding in ("utf-8-sig", "cp1256", "utf-8"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("تعذر قراءة ملف CSV؛ احفظه بترميز UTF-8.")
        delimiter = ";" if text.count(";") > text.count(",") else ","
        rows = list(csv.DictReader(io.StringIO(text), delimiter=delimiter))
        parsed = {classify_uploaded_sheet(filename, rows): clean_uploaded_rows(rows)}
    elif suffix == ".json":
        try:
            payload = json.loads(content.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as error:
            raise ValueError("ملف JSON غير صالح.") from error
        sheets = payload.get("sheets", payload) if isinstance(payload, dict) else {}
        parsed = {classify_uploaded_sheet(name, rows): clean_uploaded_rows(rows) for name, rows in sheets.items() if isinstance(rows, list)}
    else:
        raise ValueError("ارفع ملف Excel بصيغة XLSX أو ملف CSV أو JSON.")
    if not parsed or not any(parsed.values()):
        raise ValueError("لم نقرأ أي أصناف من الملف. تأكد أن الصف الأول يحتوي أسماء الأعمدة وأن الملف يحتوي بيانات.")
    # The uploaded file never controls the pharmacy directory or donations.
    return {"inventory": parsed.get("inventory", []), "supply": parsed.get("supply", []), "pharmacies": []}


def read_uploaded_multipart(handler):
    try:
        length = int(handler.headers.get("Content-Length", "0"))
    except ValueError as error:
        raise ValueError("تعذر قراءة حجم الملف.") from error
    if length < 1 or length > 12 * 1024 * 1024:
        raise ValueError("حجم الملف غير صالح؛ الحد الأقصى 12 ميجابايت.")
    content_type = handler.headers.get("Content-Type", "")
    match = re.search(r"boundary=(?:\"([^\"]+)\"|([^;]+))", content_type)
    if not match:
        raise ValueError("صيغة رفع الملف غير صالحة.")
    boundary = (match.group(1) or match.group(2)).encode("utf-8")
    body = handler.rfile.read(length)
    for part in body.split(b"--" + boundary):
        if b"filename=" not in part or b"\r\n\r\n" not in part:
            continue
        headers, file_content = part.split(b"\r\n\r\n", 1)
        filename_match = re.search(rb'filename="([^"]*)"', headers)
        filename = (filename_match.group(1) if filename_match else b"file.xlsx").decode("utf-8", errors="replace")
        file_content = file_content.rstrip(b"\r\n-")
        if not file_content:
            raise ValueError("الملف المرفوع فارغ.")
        return filename, file_content
    raise ValueError("لم يتم العثور على ملف في الطلب.")


class DawaeyHandler(SimpleHTTPRequestHandler):
    server_version = "DawaeyLocal/1.0"

    def end_headers(self) -> None:
        path = urlsplit(self.path).path
        if path.endswith((".html", ".css", ".js")):
            self.send_header("Cache-Control", "no-store, max-age=0, must-revalidate")
            self.send_header("Pragma", "no-cache")
        super().end_headers()

    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def send_json(self, status: int, payload: dict, cookie: str | None = None) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        if cookie:
            self.send_header("Set-Cookie", cookie)
        self.end_headers()
        self.wfile.write(body)

    def read_json(self) -> dict:
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError as error:
            raise ValueError("تعذر قراءة الطلب.") from error
        if length < 1 or length > 64_000:
            raise ValueError("حجم الطلب غير صالح.")
        try:
            payload = json.loads(self.rfile.read(length))
        except (json.JSONDecodeError, UnicodeDecodeError) as error:
            raise ValueError("بيانات النموذج غير صالحة.") from error
        if not isinstance(payload, dict):
            raise ValueError("بيانات النموذج غير صالحة.")
        return payload

    def verify_origin(self) -> bool:
        origin = self.headers.get("Origin")
        return not origin or urlsplit(origin).netloc == self.headers.get("Host")

    def get_session_token(self) -> str | None:
        cookie = SimpleCookie()
        try:
            cookie.load(self.headers.get("Cookie", ""))
        except Exception:
            return None
        morsel = cookie.get(SESSION_COOKIE)
        return morsel.value if morsel else None

    def authenticated_account(self, connection: sqlite3.Connection) -> sqlite3.Row | None:
        token = self.get_session_token()
        if not token:
            return None
        token_hash = hashlib.sha256(token.encode("ascii", errors="ignore")).hexdigest()
        return connection.execute(
            "SELECT accounts.* FROM sessions JOIN accounts ON accounts.id = sessions.account_id WHERE sessions.token_hash = ? AND sessions.expires_at > ?",
            (token_hash, int(time.time())),
        ).fetchone()

    def require_role(self, connection: sqlite3.Connection, role: str) -> sqlite3.Row | None:
        account = self.authenticated_account(connection)
        if not account or account["role"] != role:
            self.send_json(403, {"error": "هذه الخدمة متاحة للحساب المناسب فقط."})
            return None
        if role == "pharmacy":
            application = connection.execute("SELECT status FROM pharmacy_applications WHERE account_id = ?", (account["id"],)).fetchone()
            if not application or application["status"] != "approved":
                self.send_json(403, {"error": "حساب الصيدلية غير معتمد."})
                return None
        return account

    def do_GET(self) -> None:
        route = urlsplit(self.path).path
        if route in {"/", "/index.html"}:
            with closing(connect_database()) as connection:
                if not self.authenticated_account(connection):
                    self.send_response(302)
                    self.send_header("Location", "/auth.html")
                    self.send_header("Cache-Control", "no-store")
                    self.end_headers()
                    return
        if route == "/healthz":
            self.send_response(200)
            self.send_header("Content-Type", "text/plain; charset=utf-8")
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            self.wfile.write(b"ok")
            return
        if route.startswith("/data/") or route.startswith("/scripts/"):
            self.send_error(404)
            return
        if route == "/auth/google":
            self.handle_google_start()
            return
        if route == "/auth/google/callback":
            self.handle_google_callback()
            return
        if route == "/api/bootstrap":
            self.handle_bootstrap()
            return
        if route == "/api/area":
            self.handle_area()
            return
        if route == "/api/nearby-pharmacies":
            self.handle_nearby_pharmacies()
            return
        if route == "/api/session":
            self.handle_session()
            return
        if route == "/api/pharmacy/bootstrap":
            self.handle_pharmacy_bootstrap()
            return
        if route == "/api/pharmacy/upload":
            self.handle_pharmacy_upload()
            return
        if route == "/api/donations":
            self.handle_donations()
            return
        if route == "/api/saved-medicines":
            self.handle_saved_medicines()
            return
        if route == "/api/notifications":
            self.handle_notifications()
            return
        if route == "/webhooks/whatsapp":
            self.handle_whatsapp_verification()
            return
        super().do_GET()

    def do_HEAD(self) -> None:
        route = urlsplit(self.path).path
        if route.startswith("/data/") or route.startswith("/scripts/"):
            self.send_error(404)
            return
        super().do_HEAD()

    def do_POST(self) -> None:
        route = urlsplit(self.path).path
        if not route.startswith("/api/") and route != "/webhooks/whatsapp":
            self.send_error(404)
            return
        if route.startswith("/api/") and not self.verify_origin():
            self.send_json(403, {"error": "الطلب غير مسموح."})
            return
        try:
            payload = self.read_json() if route in ("/api/register", "/api/login", "/api/donations", "/api/saved-medicines", "/api/notifications/read", "/webhooks/whatsapp") else {}
            if route == "/api/register":
                self.handle_register(payload)
            elif route == "/api/login":
                self.handle_login(payload)
            elif route == "/api/donations":
                self.handle_create_donation(payload)
            elif route == "/api/saved-medicines":
                self.handle_save_medicine(payload)
            elif route == "/api/notifications/read":
                self.handle_mark_notifications_read()
            elif route == "/api/logout":
                self.handle_logout()
            elif route == "/api/pharmacy/upload":
                self.handle_pharmacy_upload()
            elif route == "/webhooks/whatsapp":
                self.handle_whatsapp_webhook(payload)
            else:
                self.send_json(404, {"error": "المسار غير موجود."})
        except ValueError as error:
            self.send_json(400, {"error": str(error)})
        except sqlite3.IntegrityError as error:
            message = "رقم التواصل مسجل بالفعل. جرّب تسجيل الدخول."
            if "pharmacy_applications.license_number" in str(error):
                message = "رقم الترخيص مسجل بالفعل."
            self.send_json(409, {"error": message})

    def do_PATCH(self) -> None:
        route = urlsplit(self.path).path
        if not route.startswith("/api/donations/"):
            self.send_error(404)
            return
        if not self.verify_origin():
            self.send_json(403, {"error": "الطلب غير مسموح."})
            return
        try:
            donation_id = route.rsplit("/", 1)[-1].strip()
            payload = self.read_json()
            self.handle_update_donation(donation_id, payload)
        except ValueError as error:
            self.send_json(400, {"error": str(error)})

    @staticmethod
    def donation_public(row: sqlite3.Row) -> dict:
        return {
            "id": row["id"],
            "medicine": row["medicine"],
            "area": row["area"],
            "quantity": row["quantity"],
            "unit": row["unit"],
            "nationalId": row["national_id"],
            "latitude": row["latitude"],
            "longitude": row["longitude"],
            "status": row["status"],
            "createdAt": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(row["created_at"])),
            "updatedAt": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(row["updated_at"])),
        }

    def handle_donations(self) -> None:
        with closing(connect_database()) as connection:
            if not self.require_role(connection, "pharmacy"):
                return
            rows = connection.execute(
                "SELECT id, medicine, area, quantity, unit, national_id, latitude, longitude, status, created_at, updated_at "
                "FROM donation_requests ORDER BY created_at DESC LIMIT 100"
            ).fetchall()
        self.send_json(200, {"requests": [self.donation_public(row) for row in rows]})

    def handle_create_donation(self, payload: dict) -> None:
        medicine = str(payload.get("medicine", "")).strip()
        area = str(payload.get("area", "")).strip()
        quantity = str(payload.get("quantity", "")).strip()
        unit = str(payload.get("unit", "")).strip()
        national_id = re.sub(r"\D", "", str(payload.get("nationalId", "")).strip())
        try:
            latitude = float(payload.get("latitude")) if payload.get("latitude") not in (None, "") else None
            longitude = float(payload.get("longitude")) if payload.get("longitude") not in (None, "") else None
        except (TypeError, ValueError):
            raise ValueError("بيانات GPS غير صالحة.")
        if latitude is not None and not -90 <= latitude <= 90 or longitude is not None and not -180 <= longitude <= 180:
            raise ValueError("بيانات GPS غير صالحة.")
        if not medicine or len(medicine) > 160:
            raise ValueError("اكتب اسم الدواء أو الكود بشكل صحيح.")
        if not area or len(area) > 160:
            raise ValueError("اكتب المنطقة بشكل صحيح.")
        if not quantity or len(quantity) > 40:
            raise ValueError("اكتب الكمية بشكل صحيح.")
        if unit not in {"علبة", "شريط", "شراب", "قلم", "أنبوبة مرهم", "أمبول", "بخاخ"}:
            raise ValueError("اختار نوع العبوة من القائمة.")
        if not re.fullmatch(r"\d{14}", national_id):
            raise ValueError("الرقم القومي لازم يكون 14 رقم.")
        now = int(time.time())
        donation_id = f"don-{secrets.token_urlsafe(12)}"
        token = self.get_session_token()
        account_id = None
        with closing(connect_database()) as connection, connection:
            if token:
                token_hash = hashlib.sha256(token.encode("ascii", errors="ignore")).hexdigest()
                account = connection.execute(
                    "SELECT account_id FROM sessions WHERE token_hash = ? AND expires_at > ?",
                    (token_hash, now),
                ).fetchone()
                account_id = account["account_id"] if account else None
            connection.execute(
                "INSERT INTO donation_requests(id, account_id, medicine, area, quantity, unit, national_id, latitude, longitude, status, created_at, updated_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'pending', ?, ?)",
                (donation_id, account_id, medicine, area, quantity, unit, national_id, latitude, longitude, now, now),
            )
            row = connection.execute(
                "SELECT id, medicine, area, quantity, unit, national_id, latitude, longitude, status, created_at, updated_at FROM donation_requests WHERE id = ?",
                (donation_id,),
            ).fetchone()
        self.send_json(201, {"request": self.donation_public(row)})

    def handle_update_donation(self, donation_id: str, payload: dict) -> None:
        status = str(payload.get("status", "")).strip()
        if status not in {"pending", "accepted", "completed", "rejected"}:
            raise ValueError("حالة طلب التبرع غير صالحة.")
        now = int(time.time())
        with closing(connect_database()) as connection, connection:
            pharmacy_account = self.require_role(connection, "pharmacy")
            if not pharmacy_account:
                return
            previous = connection.execute("SELECT account_id, status, medicine FROM donation_requests WHERE id = ?", (donation_id,)).fetchone()
            pharmacy = connection.execute("SELECT pharmacy_name, address, district, whatsapp, latitude, longitude FROM pharmacy_applications WHERE account_id = ?", (pharmacy_account["id"],)).fetchone()
            cursor = connection.execute(
                "UPDATE donation_requests SET status = ?, updated_at = ? WHERE id = ?",
                (status, now, donation_id),
            )
            if cursor.rowcount != 1:
                self.send_json(404, {"error": "طلب التبرع غير موجود."})
                return
            if previous and previous["account_id"] and status == "accepted" and previous["status"] != "accepted":
                connection.execute(
                    "INSERT INTO notifications(account_id, kind, title, message, related_id, created_at) VALUES (?, 'donation_accepted', ?, ?, ?, ?)",
                    (previous["account_id"], "تم قبول تبرعك", self.donation_acceptance_message(previous["medicine"], pharmacy), donation_id, now),
                )
            row = connection.execute(
                "SELECT id, medicine, area, quantity, unit, national_id, latitude, longitude, status, created_at, updated_at FROM donation_requests WHERE id = ?",
                (donation_id,),
            ).fetchone()
        self.send_json(200, {"request": self.donation_public(row)})

    @staticmethod
    def donation_acceptance_message(medicine: str, pharmacy: sqlite3.Row | None) -> str:
        if not pharmacy:
            return f"وافقت صيدلية على استلام تبرعك بدواء {medicine}. تواصل معها لتنسيق التسليم."
        location = ""
        if pharmacy["latitude"] is not None and pharmacy["longitude"] is not None:
            location = f" رابط الموقع: https://www.google.com/maps?q={pharmacy['latitude']},{pharmacy['longitude']}"
        return (f"وافقت صيدلية {pharmacy['pharmacy_name']} على استلام تبرعك بدواء {medicine}. "
                f"العنوان: {pharmacy['address']}، {pharmacy['district']}. "
                f"رقم التواصل: {pharmacy['whatsapp']}.{location}")

    def handle_saved_medicines(self) -> None:
        with closing(connect_database()) as connection:
            account = self.require_role(connection, "patient")
            if not account:
                return
            rows = connection.execute("SELECT medicine_key FROM patient_saved_medicines WHERE account_id = ? ORDER BY created_at", (account["id"],)).fetchall()
        self.send_json(200, {"medicineKeys": [row["medicine_key"] for row in rows]})

    def handle_save_medicine(self, payload: dict) -> None:
        medicine_key = str(payload.get("medicineKey", "")).strip()
        saved = bool(payload.get("saved"))
        if not medicine_key or len(medicine_key) > 120:
            raise ValueError("معرّف الدواء غير صالح.")
        with closing(connect_database()) as connection, connection:
            account = self.require_role(connection, "patient")
            if not account:
                return
            if saved:
                connection.execute("INSERT OR IGNORE INTO patient_saved_medicines(account_id, medicine_key, created_at) VALUES (?, ?, ?)", (account["id"], medicine_key, int(time.time())))
            else:
                connection.execute("DELETE FROM patient_saved_medicines WHERE account_id = ? AND medicine_key = ?", (account["id"], medicine_key))
        self.send_json(200, {"ok": True, "saved": saved, "medicineKey": medicine_key})

    def handle_notifications(self) -> None:
        with closing(connect_database()) as connection:
            account = self.require_role(connection, "patient")
            if not account:
                return
            rows = connection.execute("SELECT id, kind, title, message, related_id, is_read, created_at FROM notifications WHERE account_id = ? ORDER BY created_at DESC LIMIT 50", (account["id"],)).fetchall()
        self.send_json(200, {"notifications": [{"id": row["id"], "kind": row["kind"], "title": row["title"], "message": row["message"], "relatedId": row["related_id"], "isRead": bool(row["is_read"]), "createdAt": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(row["created_at"]))} for row in rows]})

    def handle_mark_notifications_read(self) -> None:
        with closing(connect_database()) as connection, connection:
            account = self.require_role(connection, "patient")
            if not account:
                return
            connection.execute("UPDATE notifications SET is_read = 1 WHERE account_id = ?", (account["id"],))
        self.send_json(200, {"ok": True})

    def redirect_to_auth(self, error: str) -> None:
        self.send_response(302)
        self.send_header("Location", f"/auth.html?error={urlencode({'message': error})[8:]}")
        self.end_headers()

    def handle_google_start(self) -> None:
        client_id = os.environ.get("GOOGLE_CLIENT_ID", "").strip()
        if not client_id:
            self.redirect_to_auth("تسجيل Google غير مفعّل بعد؛ أضف إعدادات Google OAuth في الخادم.")
            return
        redirect_uri = google_redirect_uri(self)
        state = secrets.token_urlsafe(32)
        query = urlencode({"client_id": client_id, "redirect_uri": redirect_uri, "response_type": "code", "scope": "openid email profile", "state": state, "access_type": "online", "prompt": "select_account"})
        self.send_response(302)
        self.send_header("Location", f"https://accounts.google.com/o/oauth2/v2/auth?{query}")
        self.send_header("Set-Cookie", oauth_state_cookie(state))
        self.end_headers()

    def handle_google_callback(self) -> None:
        parameters = parse_qs(urlsplit(self.path).query)
        if parameters.get("error"):
            self.redirect_to_auth("تم إلغاء تسجيل الدخول باستخدام Google.")
            return
        code = parameters.get("code", [""])[0]
        state = parameters.get("state", [""])[0]
        expected_state = get_cookie_value(self, OAUTH_STATE_COOKIE)
        if not state or not expected_state or not hmac.compare_digest(state, expected_state):
            self.redirect_to_auth("انتهت جلسة Google الآمنة. ابدأ تسجيل الدخول من جديد.")
            return
        client_id = os.environ.get("GOOGLE_CLIENT_ID", "").strip()
        client_secret = os.environ.get("GOOGLE_CLIENT_SECRET", "").strip()
        if not code or not client_id or not client_secret:
            self.redirect_to_auth("إعدادات Google OAuth غير مكتملة.")
            return
        try:
            token_data = google_token_exchange(code, google_redirect_uri(self), client_id, client_secret)
            profile = google_user_profile(token_data["access_token"])
            email = str(profile.get("email", "")).strip().casefold()
            if not email or profile.get("email_verified") is False:
                raise ValueError("لم يتم التحقق من بريد Google.")
            name = str(profile.get("name") or profile.get("given_name") or email.split("@", 1)[0]).strip()[:100]
            with closing(connect_database()) as connection, connection:
                account = connection.execute("SELECT * FROM accounts WHERE contact_key = ?", (email,)).fetchone()
                if account and account["role"] != "patient":
                    raise ValueError("هذا البريد مرتبط بحساب صيدلية؛ استخدم دخول الصيدلية.")
                if not account:
                    salt = secrets.token_bytes(16)
                    password = secrets.token_urlsafe(32)
                    cursor = connection.execute("INSERT INTO accounts(role, full_name, contact, contact_key, password_salt, password_hash, created_at) VALUES ('patient', ?, ?, ?, ?, ?, ?)", (name, email, email, salt.hex(), password_digest(password, salt).hex(), int(time.time())))
                    account_id = cursor.lastrowid
                    connection.execute("INSERT INTO patient_profiles(account_id, governorate, district) VALUES (?, ?, ?)", (account_id, "غير محدد", "غير محدد"))
                    account = connection.execute("SELECT * FROM accounts WHERE id = ?", (account_id,)).fetchone()
                token = create_session(connection, account["id"])
                cookie = self.session_cookie(token)
            self.send_response(302)
            self.send_header("Location", "/index.html?welcome=1#top")
            self.send_header("Set-Cookie", cookie)
            self.send_header("Set-Cookie", expired_oauth_state_cookie())
            self.end_headers()
        except (OSError, KeyError, ValueError, json.JSONDecodeError) as error:
            self.log_error("Google OAuth failed: %s", error)
            self.redirect_to_auth("تعذر تسجيل الدخول باستخدام Google. تأكد من إعداد OAuth وحاول مرة أخرى.")

    def handle_pharmacy_upload(self) -> None:
        with closing(connect_database()) as connection, connection:
            account = self.require_role(connection, "pharmacy")
            if not account:
                return
            filename, content = read_uploaded_multipart(self)
            parsed = parse_uploaded_rows(content, filename)
            connection.execute(
                "INSERT INTO pharmacy_data_files(account_id, filename, inventory_json, supply_json, pharmacies_json, uploaded_at) VALUES (?, ?, ?, ?, ?, ?) ON CONFLICT(account_id) DO UPDATE SET filename=excluded.filename, inventory_json=excluded.inventory_json, supply_json=excluded.supply_json, pharmacies_json=excluded.pharmacies_json, uploaded_at=excluded.uploaded_at",
                (account["id"], Path(filename).name[:180], json.dumps(parsed["inventory"], ensure_ascii=False), json.dumps(parsed["supply"], ensure_ascii=False), json.dumps(parsed["pharmacies"], ensure_ascii=False), int(time.time())),
            )
        self.send_json(200, {"ok": True, "filename": Path(filename).name, "counts": {key: len(value) for key, value in parsed.items()}})


    def handle_pharmacy_bootstrap(self) -> None:
        token = self.get_session_token()
        if not token:
            self.send_json(401, {"error": "يجب تسجيل الدخول بحساب صيدلية معتمد."})
            return
        token_hash = hashlib.sha256(token.encode("ascii", errors="ignore")).hexdigest()
        with closing(connect_database()) as connection:
            account = connection.execute("SELECT * FROM accounts JOIN sessions ON sessions.account_id = accounts.id WHERE sessions.token_hash = ? AND sessions.expires_at > ?", (token_hash, int(time.time()))).fetchone()
            if not account or account["role"] != "pharmacy":
                self.send_json(403, {"error": "لوحة الصيدلية متاحة لحسابات الصيدليات فقط."})
                return
            application = connection.execute("SELECT status FROM pharmacy_applications WHERE account_id = ?", (account["id"],)).fetchone()
            if not application or application["status"] != "approved":
                self.send_json(403, {"error": "حساب الصيدلية لم يتم اعتماده بعد."})
                return
            uploaded = connection.execute("SELECT filename, inventory_json, supply_json, pharmacies_json, uploaded_at FROM pharmacy_data_files WHERE account_id = ?", (account["id"],)).fetchone()
        if uploaded:
            try:
                published = json.loads(DATA_PATH.read_text(encoding="utf-8"))
                directory = published.get("sheets", {}).get("بيانات الصيداليات", [])
            except (OSError, json.JSONDecodeError, AttributeError):
                directory = []
            self.send_json(200, {"source": f"المخزون والتوريد من ملف الصيدلية: {uploaded['filename']}", "uploadedAt": uploaded["uploaded_at"], "sheets": {"الاصناف والكميات": json.loads(uploaded["inventory_json"]), "طلبات التوريد": json.loads(uploaded["supply_json"]), "بيانات الصيداليات": directory}})
            return
        self.handle_bootstrap()

    def handle_bootstrap(self) -> None:
        try:
            workbook = json.loads(DATA_PATH.read_text(encoding="utf-8"))
            sheets = workbook["sheets"]
            inventory = sheets.get("الاصناف والكميات", [])
            supply = sheets.get("طلبات التوريد", [])
            pharmacies = sheets.get("بيانات الصيداليات", [])
            supply_by_number = {str(record.get("م")): record for record in supply}
            catalog = [
                {**item, **supply_by_number.get(str(item.get("م")), {})}
                for item in inventory
            ]
            self.send_json(200, {"source": workbook.get("source", ""), "catalog": catalog, "pharmacies": pharmacies, "sheets": sheets})
        except (OSError, json.JSONDecodeError, KeyError) as error:
            self.send_json(503, {"error": "بيانات المصدر غير متاحة حاليًا."})
            self.log_error("Could not load workbook data: %s", error)

    def handle_nearby_pharmacies(self) -> None:
        """Return nearby pharmacies from OpenStreetMap, including branches outside our directory."""
        global LAST_GEOCODE_REQUEST
        parameters = parse_qs(urlsplit(self.path).query)
        try:
            latitude = float(parameters.get("lat", [""])[0])
            longitude = float(parameters.get("lon", [""])[0])
        except ValueError:
            latitude = longitude = None
        area = str(parameters.get("area", [""])[0]).strip()
        if latitude is None or longitude is None:
            if not area:
                self.send_json(400, {"error": "اكتب المحافظة أو استخدم زر GPS."})
                return
            with GEOCODE_LOCK:
                wait = 1.0 - (time.time() - LAST_GEOCODE_REQUEST)
                if wait > 0:
                    time.sleep(wait)
                query = urllib.parse.quote(f"{area}, Egypt")
                request = Request(f"https://nominatim.openstreetmap.org/search?format=jsonv2&limit=1&countrycodes=eg&q={query}", headers={"User-Agent": "Dawaey/1.0 pharmacy-directory"})
                try:
                    with urlopen(request, timeout=8) as response:
                        places = json.loads(response.read().decode("utf-8"))
                    LAST_GEOCODE_REQUEST = time.time()
                except (OSError, URLError, TimeoutError, json.JSONDecodeError):
                    self.send_json(503, {"error": "تعذر تحديد المكان الآن؛ جرّب GPS أو اسم محافظة أو منطقة أوضح."})
                    return
            if not places:
                self.send_json(404, {"error": "المكان غير موجود في مصر. جرّب اسم المحافظة أو المنطقة."})
                return
            latitude, longitude = float(places[0]["lat"]), float(places[0]["lon"])
        try:
            radius = min(max(int(float(parameters.get("radius", ["15000"])[0])), 1000), 25000)
        except ValueError:
            radius = 10000
        overpass_query = f"[out:json][timeout:18];nwr[amenity=pharmacy](around:{radius},{latitude},{longitude});out center tags;"
        request = Request("https://overpass-api.de/api/interpreter", data=urllib.parse.urlencode({"data": overpass_query}).encode("utf-8"), headers={"User-Agent": "Dawaey/1.0 pharmacy-directory", "Content-Type": "application/x-www-form-urlencoded"}, method="POST")
        payload = None
        for endpoint in ("https://overpass-api.de/api/interpreter", "https://overpass.kumi.systems/api/interpreter", "https://overpass.nchc.org.tw/api/interpreter"):
            try:
                request.full_url = endpoint
                with urlopen(request, timeout=24) as response:
                    payload = json.loads(response.read().decode("utf-8"))
                break
            except (OSError, URLError, TimeoutError, json.JSONDecodeError) as error:
                print(f"Overpass provider failed: {error}", file=sys.stderr)
        if payload is None:
            # Nominatim is slower and returns fewer records, but keeps the search
            # useful when public Overpass instances are busy or unavailable.
            fallback_query = urllib.parse.quote(f"pharmacy {area or 'Egypt'} Egypt")
            fallback_url = f"https://nominatim.openstreetmap.org/search?format=jsonv2&limit=50&countrycodes=eg&addressdetails=1&accept-language=ar,en&q={fallback_query}"
            try:
                with urlopen(Request(fallback_url, headers={"User-Agent": "Dawaey/1.0 pharmacy-directory"}), timeout=20) as response:
                    fallback_places = json.loads(response.read().decode("utf-8"))
                payload = {"elements": [{"tags": {"name": place.get("name") or "صيدلية قريبة", "display_name": place.get("display_name", ""), "addr:street": place.get("display_name", ""), "phone": (place.get("extratags") or {}).get("phone", "")}, "lat": place.get("lat"), "lon": place.get("lon")} for place in fallback_places]}
            except (OSError, URLError, TimeoutError, json.JSONDecodeError) as error:
                print(f"Nominatim pharmacy fallback failed: {error}", file=sys.stderr)
                self.send_json(503, {"error": "خدمة دليل الصيدليات غير متاحة مؤقتًا؛ جرّب بعد لحظات."})
                return
        pharmacies = []
        for item in payload.get("elements", []):
            tags = item.get("tags") or {}
            name = tags.get("name:ar") or tags.get("name") or "صيدلية قريبة"
            address = " - ".join(part for part in (tags.get("addr:state"), tags.get("addr:city"), tags.get("addr:street")) if part) or area or "بالقرب من موقعك"
            phone = tags.get("phone") or tags.get("contact:phone") or "رقم غير مدرج"
            pharmacies.append({"اسم الصيدلية": name, "المنطقة / العنوان المتوقع": address, "رقم التليفون / الخط الساخن": phone, "source": "OpenStreetMap", "lat": item.get("lat") or (item.get("center") or {}).get("lat"), "lon": item.get("lon") or (item.get("center") or {}).get("lon")})
        self.send_json(200, {"pharmacies": pharmacies[:100], "source": "OpenStreetMap", "latitude": latitude, "longitude": longitude})

    def handle_area(self) -> None:
        global LAST_GEOCODE_REQUEST
        parameters = parse_qs(urlsplit(self.path).query)
        try:
            latitude = float(parameters.get("lat", [""])[0])
            longitude = float(parameters.get("lon", [""])[0])
        except ValueError:
            self.send_json(400, {"error": "موقع جغرافي غير صالح."})
            return
        if not math.isfinite(latitude) or not math.isfinite(longitude) or not -90 <= latitude <= 90 or not -180 <= longitude <= 180:
            self.send_json(400, {"error": "موقع جغرافي غير صالح."})
            return

        with GEOCODE_LOCK:
            now = time.monotonic()
            if now - LAST_GEOCODE_REQUEST < 1:
                self.send_json(429, {"error": "استنى لحظة قبل محاولة تحديد المنطقة مرة تانية."})
                return
            LAST_GEOCODE_REQUEST = now

        query = urlencode({
            "lat": latitude,
            "lon": longitude,
            "format": "jsonv2",
            "zoom": 14,
            "addressdetails": 1,
            "namedetails": 1,
            "accept-language": "ar,en",
        })
        request = Request(
            f"https://nominatim.openstreetmap.org/reverse?{query}",
            headers={"User-Agent": "DawaeyLocal/1.0", "Accept": "application/json"},
        )
        try:
            with urlopen(request, timeout=7) as response:
                place = json.loads(response.read().decode("utf-8"))
        except (OSError, URLError, TimeoutError, json.JSONDecodeError):
            self.send_json(503, {"error": "خدمة تحديد المنطقة غير متاحة؛ اكتب منطقتك يدويًا."})
            return

        address = place.get("address", {})
        # Prefer the most useful local area, then fall back to city/governorate.
        area = next((address.get(key) for key in ("neighbourhood", "suburb", "city_district", "district", "town", "village", "hamlet", "city", "municipality", "county", "state") if address.get(key)), None)
        governorate = address.get("state") or address.get("state_district") or address.get("county") or ""
        city = address.get("city") or address.get("town") or address.get("municipality") or address.get("village") or ""
        if not area:
            self.send_json(404, {"error": "ما قدرناش نحدد اسم المنطقة من موقعك."})
            return
        self.send_json(200, {"area": str(area), "city": str(city), "governorate": str(governorate), "displayName": place.get("display_name", str(area)), "latitude": latitude, "longitude": longitude, "attribution": "OpenStreetMap contributors"})

    def handle_whatsapp_verification(self) -> None:
        parameters = parse_qs(urlsplit(self.path).query)
        mode = parameters.get("hub.mode", [""])[0]
        verify_token = parameters.get("hub.verify_token", [""])[0]
        challenge = parameters.get("hub.challenge", [""])[0]
        expected = os.environ.get("WHATSAPP_VERIFY_TOKEN", "")
        if mode == "subscribe" and expected and verify_token == expected:
            self.send_response(200); self.end_headers(); self.wfile.write(challenge.encode("utf-8")); return
        self.send_error(403)

    def handle_whatsapp_webhook(self, payload: dict) -> None:
        for entry in payload.get("entry", []):
            for change in entry.get("changes", []):
                for message in (change.get("value", {}).get("messages", []) or []):
                    button_id = ((message.get("interactive") or {}).get("button_reply") or {}).get("id", "")
                    status = handle_approval_callback(button_id)
                    if status:
                        self.send_json(200, {"ok": True, "status": status}); return
        self.send_json(200, {"ok": True})

    def handle_register(self, payload: dict) -> None:
        role = str(payload.get("role", ""))
        if role not in {"patient", "pharmacy"}:
            raise ValueError("اختار نوع الحساب.")
        full_name = str(payload.get("fullName", "")).strip()
        if len(full_name) < 2 or len(full_name) > 100:
            raise ValueError("اكتب الاسم بشكل صحيح.")
        contact, contact_key = normalize_contact(payload.get("contact"))
        password = str(payload.get("password", ""))
        if not 8 <= len(password) <= 128:
            raise ValueError("كلمة المرور لازم تكون ٨ أحرف على الأقل.")
        if payload.get("privacyAccepted") is not True:
            raise ValueError("وافق على سياسة الخصوصية لإكمال التسجيل.")

        salt = secrets.token_bytes(16)
        encoded_salt = salt.hex()
        encoded_hash = password_digest(password, salt).hex()
        with closing(connect_database()) as connection, connection:
            cursor = connection.execute(
                "INSERT INTO accounts(role, full_name, contact, contact_key, password_salt, password_hash, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (role, full_name, contact, contact_key, encoded_salt, encoded_hash, int(time.time())),
            )
            account_id = cursor.lastrowid
            if role == "patient":
                governorate = str(payload.get("governorate", "")).strip()
                district = str(payload.get("district", "")).strip()
                if not governorate or not district:
                    raise ValueError("اختار المحافظة والمنطقة.")
                connection.execute(
                    "INSERT INTO patient_profiles(account_id, governorate, district) VALUES (?, ?, ?)",
                    (account_id, governorate, district),
                )
                token = create_session(connection, account_id)
                account = connection.execute("SELECT * FROM accounts WHERE id = ?", (account_id,)).fetchone()
                cookie = self.session_cookie(token)
                self.send_json(201, {"user": public_account(connection, account)}, cookie)
                return

            pharmacy_name = str(payload.get("pharmacyName", "")).strip()
            pharmacist_name = str(payload.get("pharmacistName", "")).strip()
            license_number = str(payload.get("licenseNumber", "")).strip()
            address = str(payload.get("address", "")).strip()
            district = str(payload.get("district", "")).strip()
            opening_hours = str(payload.get("openingHours", "")).strip()
            whatsapp = str(payload.get("whatsapp", "")).strip()
            location_url = str(payload.get("locationUrl", "")).strip()
            try:
                latitude = float(payload.get("latitude")) if payload.get("latitude") not in (None, "") else None
                longitude = float(payload.get("longitude")) if payload.get("longitude") not in (None, "") else None
            except (TypeError, ValueError):
                raise ValueError("بيانات GPS للصيدلية غير صالحة.")
            if not all((pharmacy_name, pharmacist_name, license_number, address, district, opening_hours, whatsapp)):
                raise ValueError("كمّل بيانات الصيدلية والصيدلي المسؤول والعنوان ومواعيد العمل.")
            if len(license_number) > 64 or len(pharmacy_name) > 120 or len(address) > 240:
                raise ValueError("راجع طول بيانات الصيدلية.")
            connection.execute(
                "INSERT INTO pharmacy_applications(account_id, pharmacy_name, pharmacist_name, license_number, address, district, opening_hours, whatsapp, latitude, longitude, location_url) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (account_id, pharmacy_name, pharmacist_name, license_number, address, district, opening_hours, whatsapp, latitude, longitude, location_url),
            )
        telegram_sent = notify_pharmacy_application(account_id, {"pharmacyName": pharmacy_name, "pharmacistName": pharmacist_name, "licenseNumber": license_number, "address": address, "district": district, "contact": contact, "locationUrl": location_url})
        message = "تم حفظ طلب الصيدلية وإرسال إشعار المراجعة." if telegram_sent else "تم حفظ طلب الصيدلية، لكن تعذر إرسال إشعار تيليجرام. راجع إعدادات Telegram في Railway."
        self.send_json(202, {"status": "pending", "telegramSent": telegram_sent, "message": message})

    def handle_login(self, payload: dict) -> None:
        _, contact_key = normalize_contact(payload.get("contact"))
        password = str(payload.get("password", ""))
        role = str(payload.get("role", ""))
        with closing(connect_database()) as connection, connection:
            account = connection.execute("SELECT * FROM accounts WHERE contact_key = ?", (contact_key,)).fetchone()
            if not account or not hmac.compare_digest(
                bytes.fromhex(account["password_hash"]),
                password_digest(password, bytes.fromhex(account["password_salt"])),
            ):
                self.send_json(401, {"error": "بيانات الدخول مش صحيحة."})
                return
            if role and account["role"] != role:
                self.send_json(401, {"error": "نوع الحساب لا يطابق بيانات الدخول."})
                return
            if account["role"] == "pharmacy":
                application = connection.execute(
                    "SELECT status FROM pharmacy_applications WHERE account_id = ?", (account["id"],)
                ).fetchone()
                if not application or application["status"] != "approved":
                    self.send_json(403, {"error": "طلب الصيدلية قيد المراجعة. هيتفعل الحساب بعد اعتماد البيانات."})
                    return
            token = create_session(connection, account["id"])
            self.send_json(200, {"user": public_account(connection, account)}, self.session_cookie(token))

    def handle_session(self) -> None:
        token = self.get_session_token()
        if not token:
            self.send_json(200, {"user": None})
            return
        token_hash = hashlib.sha256(token.encode("ascii", errors="ignore")).hexdigest()
        with closing(connect_database()) as connection, connection:
            account = connection.execute(
                "SELECT accounts.* FROM sessions JOIN accounts ON accounts.id = sessions.account_id WHERE sessions.token_hash = ? AND sessions.expires_at > ?",
                (token_hash, int(time.time())),
            ).fetchone()
            if not account:
                connection.execute("DELETE FROM sessions WHERE token_hash = ?", (token_hash,))
                self.send_json(200, {"user": None}, self.expired_cookie())
                return
            self.send_json(200, {"user": public_account(connection, account)})

    def handle_logout(self) -> None:
        token = self.get_session_token()
        if token:
            token_hash = hashlib.sha256(token.encode("ascii", errors="ignore")).hexdigest()
            with closing(connect_database()) as connection, connection:
                connection.execute("DELETE FROM sessions WHERE token_hash = ?", (token_hash,))
        self.send_json(200, {"ok": True}, self.expired_cookie())

    @staticmethod
    def session_cookie(token: str) -> str:
        secure = os.environ.get("DAWAEY_SECURE_COOKIE") == "1"
        same_site = "None" if secure else "Lax"
        suffix = "; Secure" if secure else ""
        return f"{SESSION_COOKIE}={token}; HttpOnly; SameSite={same_site}; Path=/; Max-Age={SESSION_SECONDS}{suffix}"

    @staticmethod
    def expired_cookie() -> str:
        secure = os.environ.get("DAWAEY_SECURE_COOKIE") == "1"
        same_site = "None" if secure else "Lax"
        suffix = "; Secure" if secure else ""
        return f"{SESSION_COOKIE}=; HttpOnly; SameSite={same_site}; Path=/; Max-Age=0{suffix}"


def main() -> None:
    if not DATA_PATH.exists():
        raise FileNotFoundError(f"Missing published dataset: {DATA_PATH}")
    print("Dawaey published dataset loaded")
    initialize_database()
    start_telegram_polling()
    port = int(os.environ.get("PORT", "5173"))
    server = ThreadingHTTPServer(("", port), DawaeyHandler)
    print(f"Dawaey server listening on http://localhost:{port}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nDawaey server stopped")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
